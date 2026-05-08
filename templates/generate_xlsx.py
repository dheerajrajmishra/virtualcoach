"""Run this script once to generate the three Excel template .xlsx files.
Requires: pip install openpyxl
"""
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

HERE = Path(__file__).parent
HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(color="FFFFFF", bold=True)

DROPDOWNS = {
    "transcripts_template.xlsx": {},
    "faqs_template.xlsx": {
        7: '"all,en,hi,ta,te,mr,bn"'   # language_scope column index (1-based)
    },
    "quizzes_template.xlsx": {
        4: '"text,audio,video"',        # input_type
        8: '"all,en,hi,ta,te,mr,bn"'   # language_scope
    },
}


def csv_to_xlsx(csv_name: str, xlsx_name: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    csv_path = HERE / csv_name
    with open(csv_path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, start=1):
            for col_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if row_idx == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.alignment = Alignment(horizontal="center")

    # Auto-size columns
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

    # Add data validation dropdowns
    from openpyxl.worksheet.datavalidation import DataValidation
    for col_idx, formula in DROPDOWNS.get(xlsx_name, {}).items():
        col_letter = get_column_letter(col_idx)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.sqref = f"{col_letter}2:{col_letter}1000"
        ws.add_data_validation(dv)

    wb.save(HERE / xlsx_name)
    print(f"Created {xlsx_name}")


if __name__ == "__main__":
    csv_to_xlsx("transcripts_template.csv", "transcripts_template.xlsx")
    csv_to_xlsx("faqs_template.csv", "faqs_template.xlsx")
    csv_to_xlsx("quizzes_template.csv", "quizzes_template.xlsx")
    print("Done. All three Excel templates generated.")
